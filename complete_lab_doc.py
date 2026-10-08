from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\ander\Downloads\GLAB-S08-EAREVALO-2026-2.docx")
ROOT = Path(r"C:\Users\ander\source\repos\lab08")
CAPTURES = ROOT / "capturas"
OUTPUT = ROOT / "GLAB-S08-EAREVALO-2026-2-COMPLETADO.docx"


def set_run_font(run, name="Arial", size=10.5, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def format_paragraph(paragraph, *, before=0, after=6, line=1.08):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_heading(doc, text, level=1, page_break=False):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 12, bold=True)
    format_paragraph(paragraph, before=8 if level == 1 else 6, after=5)
    return paragraph


def add_body(doc, text, *, bold_lead=None, after=6):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    format_paragraph(paragraph, after=after)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
    run = paragraph.add_run("• " + text)
    set_run_font(run)
    format_paragraph(paragraph, after=3)
    return paragraph


NUMBER_COUNTER = 0


def add_numbered(doc, text):
    global NUMBER_COUNTER
    NUMBER_COUNTER += 1
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    run = paragraph.add_run(f"{NUMBER_COUNTER}. {text}")
    set_run_font(run)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    format_paragraph(paragraph, after=4)
    return paragraph


def add_capture(doc, filename, caption, alt_text):
    image_path = CAPTURES / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(2)
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", caption)

    caption_paragraph = doc.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(7)
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=9, italic=True)


doc = Document(SOURCE)

# Some institutional templates omit standard built-in styles.
for missing_style in ("Caption", "List Bullet", "List Number"):
    try:
        doc.styles[missing_style]
    except KeyError:
        doc.styles.add_style(missing_style, WD_STYLE_TYPE.PARAGRAPH)

# Configure only the styles used by the appended delivery section.
for style_name, size, bold in (("Heading 1", 16, True), ("Heading 2", 12, True)):
    style = doc.styles[style_name]
    style.font.name = "Arial"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)

caption_style = doc.styles["Caption"]
caption_style.font.name = "Arial"
caption_style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
caption_style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
caption_style.font.size = Pt(9)
caption_style.font.color.rgb = RGBColor(0, 0, 0)

# Replace the original blank delivery prompts with a completed submission heading.
delivery = doc.paragraphs[39]
delivery.clear()
delivery.style = doc.styles["Heading 1"]
delivery.paragraph_format.keep_with_next = True
delivery_run = delivery.add_run("Entrega del laboratorio con NeptunoDB")
set_run_font(delivery_run, size=16, bold=True)
format_paragraph(delivery, before=0, after=5)

add_body(
    doc,
    "La solución adapta el ejercicio MVC con Dapper a la base de datos NeptunoDB ya existente. "
    "El mantenimiento principal se realizó sobre Productos, el registro auxiliar sobre Proveedores "
    "y el reporte sobre Pedidos y DetallePedidos.",
)

add_heading(doc, "Repositorio y archivos entregados", level=2)
add_bullet(doc, "Repositorio de referencia: https://github.com/EdwinArevalo/GradesMVC.git")
add_bullet(doc, "Solución desarrollada: lab08.slnx y proyecto lab08/lab08.csproj.")
add_bullet(doc, "Script de base de datos: lab08/Script-NeptunoDB.sql.")
add_bullet(doc, "Base utilizada: NeptunoDB en SQL Server, sin crear una base adicional.")

add_heading(doc, "Capturas del recorrido de la aplicación", level=1, page_break=True)
add_capture(
    doc,
    "01-inicio.jpg",
    "Figura 1. Página de inicio y menú principal.",
    "Página de inicio de Neptuno MVC con enlaces a productos y reporte de pedidos.",
)
add_capture(
    doc,
    "02-productos.jpg",
    "Figura 2. Listado de productos activos con acciones CRUD.",
    "Listado de productos activos con categoría, proveedor, precio, stock y acciones.",
)

add_heading(doc, "Búsqueda y consulta de productos", level=2, page_break=True)
add_capture(
    doc,
    "03-busqueda-producto.jpg",
    "Figura 3. Búsqueda por nombre enviada mediante GET.",
    "Resultado de la búsqueda del producto Salsa de Ají Amarillo.",
)
add_capture(
    doc,
    "04-detalle-producto.jpg",
    "Figura 4. Detalle del producto seleccionado.",
    "Vista de detalle con categoría, proveedor, precio, existencias y estado.",
)

add_heading(doc, "Edición y proveedores", level=2, page_break=True)
add_capture(
    doc,
    "05-editar-producto.jpg",
    "Figura 5. Formulario de edición con listas de proveedor y categoría.",
    "Formulario para editar un producto con listas desplegables y campos validados.",
)
add_capture(
    doc,
    "06-proveedores.jpg",
    "Figura 6. Listado de proveedores activos.",
    "Listado de proveedores activos con información de contacto.",
)

add_heading(doc, "Reporte y recorrido completo de una petición", level=1, page_break=True)
add_capture(
    doc,
    "07-reporte-pedidos.jpg",
    "Figura 7. Reporte de pedidos filtrado del 1 al 31 de agosto de 2026.",
    "Reporte de pedidos con cliente, producto, cantidad, subtotal, fecha requerida y estado.",
)

add_heading(doc, "Recorrido de la búsqueda de productos", level=2)
add_numbered(doc, "Ruta: el navegador envía GET /Productos?busqueda=Salsa mediante el formulario de búsqueda.")
add_numbered(doc, "Controlador: ProductosController.Index recibe el parámetro busqueda, lo conserva en ViewData[\"Busqueda\"] y solicita la información al repositorio.")
add_numbered(doc, "Repositorio: ProductoRepositorio.ListarAsync abre la conexión definida en appsettings.json y ejecuta QueryAsync<Producto> con CommandType.StoredProcedure.")
add_numbered(doc, "Procedimiento almacenado: usp_Lab08_Producto_Listar filtra por nombre, considera únicamente registros con Activo = 1 y relaciona Productos con Proveedores y Categorias.")
add_numbered(doc, "Modelo: Dapper mapea el resultado a IEnumerable<Producto>, incluyendo NombreProveedor y NombreCategoria obtenidos mediante JOIN.")
add_numbered(doc, "Vista: Views/Productos/Index.cshtml recibe el modelo fuertemente tipado y reutiliza _ProductoRow.cshtml para representar cada fila.")

add_heading(doc, "Datos enviados a las vistas", level=2, page_break=True)
add_body(
    doc,
    "Modelo: se utiliza para transportar los productos, proveedores o filas del reporte porque constituyen el contenido principal y requieren tipado fuerte.",
    bold_lead="Modelo:",
)
add_body(
    doc,
    "ViewData: conserva el texto de búsqueda, el intervalo de fechas y las listas desplegables de proveedores y categorías. Estos valores complementan al modelo principal.",
    bold_lead="ViewData:",
)
add_body(
    doc,
    "TempData: muestra mensajes de confirmación después de crear, editar o eliminar. Se eligió porque el mensaje debe sobrevivir a la redirección del patrón Post/Redirect/Get.",
    bold_lead="TempData:",
)

add_heading(doc, "Flujos adicionales verificados", level=2)
add_bullet(doc, "Create y Edit validan DataAnnotations y vuelven a cargar los desplegables cuando ModelState no es válido.")
add_bullet(doc, "Delete ejecuta una eliminación lógica mediante Activo = 0; no se realiza DELETE físico.")
add_bullet(doc, "Proveedores valida el nombre de compañía duplicado con ModelState.AddModelError.")
add_bullet(doc, "PedidosController.Reporte conserva desde y hasta en ViewData y consulta por GET para mantener el filtro visible.")
add_bullet(doc, "Los controladores no contienen SQL ni abren conexiones; todas las operaciones de datos utilizan async/await.")

add_heading(doc, "Observaciones y conclusiones", level=2)
add_body(
    doc,
    "La separación entre controladores, repositorios, procedimientos almacenados y vistas permite mantener responsabilidades claras. "
    "Dapper concentra el acceso a datos en métodos asíncronos, mientras que MVC gestiona validación, navegación y presentación.",
)
add_body(
    doc,
    "La adaptación a NeptunoDB conserva los objetivos técnicos de la guía: CRUD, búsqueda, INNER JOIN, listas desplegables, "
    "reporte por fechas, eliminación lógica, inyección de dependencias y patrón Post/Redirect/Get.",
)
add_body(
    doc,
    "La solución compiló sin errores ni advertencias. Se verificaron las páginas de inicio, productos, búsqueda, detalle, edición, "
    "proveedores y reporte; además, las pruebas transaccionales de inserción, actualización y eliminación lógica no dejaron datos residuales.",
)

# Preserve the institutional artwork and add fallback alternative text where
# the source template omitted it.
for doc_pr in doc._element.xpath(".//wp:docPr"):
    if not doc_pr.get("descr"):
        doc_pr.set("descr", "Elemento gráfico institucional de Tecsup.")
    if not doc_pr.get("title"):
        doc_pr.set("title", "Identidad visual de Tecsup")
for section in doc.sections:
    for container in (section.header, section.first_page_header, section.even_page_header):
        for doc_pr in container._element.xpath(".//wp:docPr"):
            if not doc_pr.get("descr"):
                doc_pr.set("descr", "Logotipo institucional de Tecsup.")
            if not doc_pr.get("title"):
                doc_pr.set("title", "Logotipo de Tecsup")

doc.core_properties.title = "Laboratorio 08 MVC con Dapper en NeptunoDB"
doc.core_properties.subject = "Evidencias, recorrido de la aplicación y conclusiones"
doc.save(OUTPUT)
print(OUTPUT)
